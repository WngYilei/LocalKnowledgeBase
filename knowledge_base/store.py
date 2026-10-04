"""Document ingestion and Chroma retrieval for the standalone knowledge base."""

from __future__ import annotations

import hashlib
import io
import json
import os
import re
import sys
import tempfile
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .embeddings import get_embedding_provider
from .sanitizer import SANITIZER_VERSION, contains_common_pii, desensitize_text, sanitize_text

try:
    import pysqlite3 as sqlite3  # type: ignore

    sys.modules["sqlite3"] = sqlite3
except Exception:  # pragma: no cover
    import sqlite3  # type: ignore


PROJECT_DIR = Path(__file__).resolve().parents[1]
DATA_DIR = Path(os.getenv("KNOWLEDGE_BASE_DATA_DIR", str(PROJECT_DIR / "data"))).expanduser().resolve()
UPLOAD_DIR = DATA_DIR / "uploads"
CATALOG_PATH = DATA_DIR / "catalog.json"
COLLECTION_NAME = os.getenv("KNOWLEDGE_BASE_COLLECTION", "knowledge_hub_v1")
CHUNK_SIZE = 900
CHUNK_OVERLAP = 120
LOCK = threading.Lock()
INIT_LOCK = threading.Lock()
INITIALIZED = False

DEMO_DOCUMENTS = [
    ("seed-1", "Chroma 是一个开源向量数据库，适合做语义检索和 RAG。"),
    ("seed-2", "Embedding 模型会把文本转换成向量，语义相近的文本向量更接近。"),
    ("seed-3", "RAG 通常先检索知识库，再把检索结果交给大模型生成答案。"),
    ("seed-4", "本地持久化目录可以让 Chroma 在重启后继续保留已写入的数据。"),
]

def ensure_storage() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    if not CATALOG_PATH.exists():
        CATALOG_PATH.write_text("[]", encoding="utf-8")


def _chromadb():
    import chromadb

    return chromadb


def _seed_demo_enabled() -> bool:
    return os.getenv("KNOWLEDGE_BASE_SEED_DEMO", "true").strip().lower() in {"1", "true", "yes", "on"}


def open_collection():
    ensure_storage()
    client = _chromadb().PersistentClient(path=str(DATA_DIR))
    return client.get_or_create_collection(name=COLLECTION_NAME, metadata={"hnsw:space": "cosine"})


def ensure_initialized() -> None:
    global INITIALIZED
    if INITIALIZED:
        return
    with INIT_LOCK:
        if INITIALIZED:
            return
        collection = open_collection()
        if collection.count() == 0 and _seed_demo_enabled():
            provider = get_embedding_provider()
            collection.upsert(
                ids=[item[0] for item in DEMO_DOCUMENTS],
                documents=[item[1] for item in DEMO_DOCUMENTS],
                metadatas=[
                    {
                        "source_kind": "seed",
                        "file_name": "demo",
                        "page_number": 1,
                        "sanitizer_version": SANITIZER_VERSION,
                        "embedding_model": provider.model_name,
                    }
                    for _ in DEMO_DOCUMENTS
                ],
                embeddings=provider.embed_documents([item[1] for item in DEMO_DOCUMENTS]),
            )
        INITIALIZED = True


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def clean_markdown_text(text: str) -> str:
    text = re.sub(r"```.*?```", lambda match: match.group(0).strip("`"), text, flags=re.S)
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"^\s*[-*+]\s+", "", text, flags=re.M)
    text = re.sub(r"^\s*\d+\.\s+", "", text, flags=re.M)
    return normalize_text(re.sub(r"[*_]{1,3}", "", text))


def sanitize_filename(filename: str) -> str:
    name = Path(filename).name
    name = re.sub(r"[\x00-\x1f\x7f/\\]+", "_", name)
    name = re.sub(r"[^\w.\-]+", "_", name, flags=re.UNICODE)
    return name.strip("._") or "upload.bin"


def _extract_segments(filename: str, content_type: str | None, data: bytes) -> tuple[str, list[dict[str, Any]]]:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf" or content_type == "application/pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
        segments = []
        for page_number, page in enumerate(reader.pages, 1):
            text = normalize_text(page.extract_text() or "")
            if text:
                segments.append({"page_number": page_number, "text": text})
        return "pdf", segments

    if ext == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        from docx import Document

        document = Document(io.BytesIO(data))
        parts = [normalize_text(paragraph.text) for paragraph in document.paragraphs if normalize_text(paragraph.text)]
        for table_index, table in enumerate(document.tables, 1):
            rows = [" | ".join(normalize_text(cell.text) for cell in row.cells if normalize_text(cell.text)) for row in table.rows]
            rows = [row for row in rows if row]
            if rows:
                parts.append(f"[Table {table_index}]\n" + "\n".join(rows))
        text = normalize_text("\n\n".join(parts))
        return "docx", [{"page_number": 1, "text": text}] if text else []

    if ext in {".md", ".markdown", ".txt"} or (content_type or "").startswith("text/"):
        text = data.decode("utf-8", errors="ignore")
        text = clean_markdown_text(text) if ext in {".md", ".markdown"} else normalize_text(text)
        return ("markdown" if ext in {".md", ".markdown"} else "text"), ([{"page_number": 1, "text": text}] if text else [])

    raise ValueError("Only PDF, DOCX, Markdown, and TXT files are supported.")


def chunk_text(text: str) -> list[str]:
    text = normalize_text(text)
    if len(text) <= CHUNK_SIZE:
        return [text] if text else []
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + CHUNK_SIZE)
        if text[start:end].strip():
            chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(0, end - CHUNK_OVERLAP)
    return chunks


def _load_catalog() -> list[dict[str, Any]]:
    ensure_storage()
    try:
        value = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
        return value if isinstance(value, list) else []
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Catalog is corrupted: {CATALOG_PATH}") from exc


def _save_catalog(items: list[dict[str, Any]]) -> None:
    ensure_storage()
    with tempfile.NamedTemporaryFile(
        mode="w", encoding="utf-8", dir=DATA_DIR, prefix="catalog.", suffix=".tmp", delete=False
    ) as handle:
        temporary_path = Path(handle.name)
        json.dump(items, handle, ensure_ascii=False, indent=2)
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temporary_path, CATALOG_PATH)


def _public_file_item(item: dict[str, Any]) -> dict[str, Any]:
    """Return only metadata safe for Web/MCP/LLM callers."""
    raw_name = str(item.get("file_name") or item.get("stored_name") or "document")
    safe_name = sanitize_filename(desensitize_text(raw_name))
    return {
        "file_hash": str(item.get("file_hash", "")),
        "file_name": safe_name,
        "source_kind": str(item.get("source_kind", "")),
        "size_bytes": int(item.get("size_bytes", 0) or 0),
        "uploaded_at": item.get("uploaded_at"),
        "segment_count": int(item.get("segment_count", 0) or 0),
        "chunk_count": int(item.get("chunk_count", 0) or 0),
    }


def ingest_file(filename: str, content_type: str | None, data: bytes) -> dict[str, Any]:
    ensure_initialized()
    ensure_storage()
    source_kind, segments = _extract_segments(filename, content_type, data)
    segments = [{**segment, "text": sanitize_text(segment["text"])[0]} for segment in segments]
    records = []
    file_hash = hashlib.sha256(data).hexdigest()
    safe_name = sanitize_filename(filename)
    stored_name = f"{file_hash[:12]}_{safe_name}"
    stored_path = UPLOAD_DIR / stored_name

    for segment_index, segment in enumerate(segments):
        chunks = chunk_text(segment["text"])
        for chunk_index, chunk in enumerate(chunks):
            records.append((
                f"{file_hash}:{segment_index}:{chunk_index}",
                chunk,
                {
                    "source_kind": source_kind, "file_name": filename, "stored_name": stored_name,
                    "file_hash": file_hash,
                    "content_type": content_type or "", "segment_index": segment_index,
                    "segment_count": len(segments), "chunk_index": chunk_index,
                    "chunk_count": len(chunks), "page_number": segment.get("page_number"),
                    "uploaded_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                    "size_bytes": len(data),
                    "sanitizer_version": SANITIZER_VERSION,
                },
            ))
    if not records:
        raise ValueError("No extractable text was found in this file.")

    provider = get_embedding_provider()
    for record in records:
        record[2]["embedding_model"] = provider.model_name
        record[2]["embedding_dimension"] = provider.dimension
    embeddings = provider.embed_documents([item[1] for item in records])

    with LOCK:
        stored_path.write_bytes(data)
        collection = open_collection()
        collection.upsert(
            ids=[item[0] for item in records],
            documents=[item[1] for item in records],
            metadatas=[item[2] for item in records],
            embeddings=embeddings,
        )
        catalog = [item for item in _load_catalog() if item.get("file_hash") != file_hash]
        catalog.insert(0, {
            "file_hash": file_hash, "file_name": filename, "stored_name": stored_name,
            "stored_path": str(Path("uploads") / stored_name), "content_type": content_type or "",
            "source_kind": source_kind, "size_bytes": len(data),
            "uploaded_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "segment_count": len(segments), "chunk_count": len(records),
            "sanitizer_version": SANITIZER_VERSION,
            "embedding_model": provider.model_name,
            "embedding_dimension": provider.dimension,
        })
        _save_catalog(catalog)

    result = {
        "file_hash": file_hash,
        "file_name": filename,
        "source_kind": source_kind,
        "segment_count": len(segments),
        "chunk_count": len(records),
        "sanitizer_version": SANITIZER_VERSION,
        "embedding_model": provider.model_name,
        "embedding_dimension": provider.dimension,
    }
    if os.getenv("KNOWLEDGE_BASE_ENABLE_RAW_DOWNLOAD", "false").strip().lower() in {"1", "true", "yes", "on"}:
        result["download_url"] = f"/api/files/{file_hash}"
    return result


def list_files(limit: int = 20) -> list[dict[str, Any]]:
    ensure_initialized()
    return [_public_file_item(item) for item in _load_catalog()[:limit]]


def _safe_result_metadata(metadata: dict[str, Any]) -> dict[str, Any]:
    file_hash = str(metadata.get("file_hash", ""))
    raw_name = str(metadata.get("file_name", ""))
    return {
        "source_id": file_hash,
        "file_name": sanitize_filename(desensitize_text(raw_name)) if raw_name else "demo",
        "source_kind": str(metadata.get("source_kind", "")),
        "page_number": metadata.get("page_number"),
        "segment_index": int(metadata.get("segment_index", 0) or 0),
        "chunk_index": int(metadata.get("chunk_index", 0) or 0),
    }


def sanitize_search_results(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Apply the final redaction and output schema before data leaves the module."""
    results: list[dict[str, Any]] = []
    for row in rows:
        content = sanitize_text(str(row.get("content", "")))[0]
        metadata = _safe_result_metadata(row.get("metadata") or {})
        if contains_common_pii(content) or contains_common_pii(metadata["file_name"]):
            raise RuntimeError("Search result failed the final sanitization check")
        result = {
            "id": str(row.get("id", "")),
            "content": content,
            "snippet": content[:360] + ("..." if len(content) > 360 else ""),
            "metadata": metadata,
            "distance": float(row.get("distance", 0.0)),
            "file_name": metadata["file_name"],
            "source_kind": metadata["source_kind"],
            "page_number": metadata["page_number"],
            "chunk_index": metadata["chunk_index"],
        }
        results.append(result)
    return results


def search_documents(query: str, top_k: int = 5) -> list[dict[str, Any]]:
    ensure_initialized()
    query = desensitize_text(query.strip())
    if not query:
        return []
    if not 1 <= top_k <= 50:
        raise ValueError("top_k must be between 1 and 50")
    provider = get_embedding_provider()
    result = open_collection().query(
        query_embeddings=[provider.embed_query(query)],
        n_results=top_k,
        include=["documents", "metadatas", "distances"],
    )
    rows = [
        {"id": doc_id, "content": document, "metadata": metadata, "distance": distance}
        for doc_id, document, metadata, distance in zip(
            result["ids"][0], result["documents"][0], result["metadatas"][0], result["distances"][0]
        )
    ]
    return sanitize_search_results(rows)


def build_state() -> dict[str, Any]:
    ensure_initialized()
    return {
        "collection": COLLECTION_NAME,
        "document_count": open_collection().count(),
        "file_count": len(_load_catalog()),
        "files": list_files(),
    }


def get_file_path(file_hash: str) -> tuple[Path, str] | None:
    if not re.fullmatch(r"[0-9a-f]{64}", file_hash):
        return None
    for item in _load_catalog():
        if item.get("file_hash") == file_hash:
            stored_name = Path(str(item.get("stored_name", ""))).name
            path = (UPLOAD_DIR / stored_name).resolve()
            upload_root = UPLOAD_DIR.resolve()
            if upload_root not in path.parents:
                return None
            return path, sanitize_filename(desensitize_text(str(item.get("file_name", path.name))))
    return None
